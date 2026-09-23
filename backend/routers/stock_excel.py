"""
WMS Kiwkiw - Lançamento de movimentações de estoque por planilha (23/09/2026)

Tela de Estoque (interna) → botão "Lançar por Excel". Serve para acerto de
inventário, entrada que chegou sem NF e casos parecidos, que antes eram
lançados um a um pelo "Lançamento Manual".

Fluxo em 3 passos, mesmo desenho de Devoluções:

  GET  /inventory/planilha/modelo               -> Excel modelo em memória
  POST /inventory/planilha/{seller_id}/analyze  -> confere, NÃO grava nada
  POST /inventory/planilha/{seller_id}/lancar   -> revalida tudo e grava

REGRAS (decididas com o dono do sistema em 23/09/2026):

  * UM arquivo por seller — o selecionado no topo da tela. A planilha não
    tem coluna Seller.
  * A data vem da planilha (pode ser retroativa). Vazia, inválida ou FUTURA
    bloqueia — nunca vira "hoje" sozinha.
  * TUDO-OU-NADA: qualquer linha com erro bloqueia o lote inteiro.
  * Bloqueiam: data, tipo (Entrada/Saída/E/S), SKU sem produto ATIVO no
    seller, SKU descontinuado, quantidade que não seja inteiro > 0, e NF e
    Observação as duas vazias (pelo menos uma é obrigatória).
  * SÓ AVISAM (não bloqueiam): saldo que fica negativo, linhas idênticas no
    mesmo arquivo e linha que parece já lançada antes (mesmo seller, SKU,
    data, tipo, quantidade, NF e observação) — o jeito de pegar o mesmo
    arquivo subido duas vezes.
  * Natureza gravada: "Lançamento de planilha".
  * ⚠️ O movimento NASCE SEM `order_id`, de propósito: amarrá-lo a uma NF de
    venda faria `reverse_stock_for_order()` (saldo líquido por order_id)
    varrer este lançamento junto ao cancelar/inativar aquela NF.
  * O saldo usa `update_stock_position` (soma atômica no banco, 16/09/2026) —
    não somar em Python aqui.
  * Admin e gerente; gerente só nos sellers que atende.

Os dois endpoints antigos de lote (`/inventory/movements/bulk` e
`/inventory/bulk-stock-upload`) continuam existindo sem botão na tela e NÃO
são usados aqui — ver CLAUDE.md.
"""

import io
from collections import defaultdict
from datetime import date, datetime
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException, File, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy import func, cast, String
from sqlalchemy.orm import Session

from ..database import get_db
from ..auth import require_manager_or_above, get_user_seller_ids
from ..timezone_utils import now_brasilia, today_brasilia
from .. import models
from ..services.stock_manager import update_stock_position
from .returns import _parse_qtd, _cell_str, _strip_accents

router = APIRouter(prefix="/inventory/planilha", tags=["Estoque — Lançamento por planilha"])

NATUREZA = "Lançamento de planilha"
ABA = "MOVIMENTACOES"
COLUNAS_MODELO = ["Data", "Tipo", "SKU", "Quantidade", "NF", "Observação"]

# `stock_movements.nf_number` é VARCHAR(20): NF maior bloqueia em vez de ser
# cortada em silêncio.
NF_MAX = 20

_ENTRADA = {"entrada", "e"}
_SAIDA = {"saida", "s"}

# Tipo gravado no banco (enum) -> rótulo devolvido à tela.
_TIPO_LABEL = {
    models.MovementType.IN: "Entrada",
    models.MovementType.OUT: "Saída",
}


# ─────────────────────────────────────────────────────────
# ESCOPO
# ─────────────────────────────────────────────────────────

def _get_seller_in_scope(
    seller_id: int, user_seller_ids: Optional[List[int]], db: Session,
) -> models.Seller:
    """Seller ativo e atendido pelo usuário (None = admin, sem filtro)."""
    if user_seller_ids is not None and seller_id not in (user_seller_ids or []):
        raise HTTPException(status_code=403, detail="Você não atende este seller")
    seller = db.query(models.Seller).filter(
        models.Seller.id == seller_id,
        models.Seller.active == True,  # noqa: E712
    ).first()
    if not seller:
        raise HTTPException(status_code=404, detail=f"Seller {seller_id} não encontrado ou inativo")
    return seller


# ─────────────────────────────────────────────────────────
# MODELO DA PLANILHA
# ─────────────────────────────────────────────────────────

@router.get("/modelo")
def download_modelo(
    current_user: models.User = Depends(require_manager_or_above),
):
    """Gera o Excel modelo EM MEMÓRIA (em produção o disco é efêmero)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = ABA

    ws.append(COLUNAS_MODELO)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7B63E8")
        cell.alignment = Alignment(horizontal="center")

    hoje = today_brasilia()
    ws.append([hoje, "Entrada", "SKU-EXEMPLO", 10, "123456", ""])
    ws.append([hoje, "Saída", "SKU-EXEMPLO", 2, "", "Avaria no armazém"])
    for row in ws.iter_rows(min_row=2, max_row=3, min_col=1, max_col=1):
        for cell in row:
            cell.number_format = "DD/MM/YYYY"

    larguras = [14, 12, 26, 13, 16, 45]
    for i, w in enumerate(larguras, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    ws2 = wb.create_sheet("INSTRUCOES")
    for linha in [
        ["Como preencher"],
        [""],
        ["", "Um arquivo por seller: ele vale para o seller selecionado na tela de Estoque."],
        [""],
        ["Data", "Data da movimentacao (dd/mm/aaaa). Obrigatoria. Pode ser no passado, nunca no futuro."],
        ["Tipo", "Entrada ou Saida (aceita tambem E ou S)."],
        ["SKU", "SKU do seller. Precisa ter produto ativo cadastrado e nao pode estar descontinuado."],
        ["Quantidade", "Numero inteiro maior que zero."],
        ["NF", "Numero da NF. Opcional, mas NF ou Observacao tem que estar preenchida."],
        ["Observacao", "Motivo do lancamento. Opcional, mas NF ou Observacao tem que estar preenchida."],
        [""],
        ["Observacoes"],
        ["", "Apague as 2 linhas de exemplo antes de subir."],
        ["", "Se qualquer linha tiver problema, NADA e lancado."],
        ["", "Saldo que fica negativo, linha repetida ou linha que parece ja lancada so geram aviso."],
    ]:
        ws2.append(linha)
    ws2.column_dimensions["A"].width = 14
    ws2.column_dimensions["B"].width = 100

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="MODELO_LANCAMENTO_ESTOQUE.xlsx"'},
    )


# ─────────────────────────────────────────────────────────
# CONVERSÃO DE CÉLULAS
# ─────────────────────────────────────────────────────────

def _parse_data(raw) -> Optional[date]:
    """
    Aceita a data como o Excel entrega (datetime/date ou número de série),
    texto dd/mm/aaaa e o ISO aaaa-mm-dd (que é como a tela devolve as linhas
    conferidas no /lancar). Qualquer outra coisa -> None (vira erro).
    """
    if raw is None:
        return None
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        try:
            from openpyxl.utils.datetime import from_excel
            val = from_excel(raw)
        except Exception:
            return None
        return val.date() if isinstance(val, datetime) else val
    txt = str(raw).strip()
    if not txt:
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(txt[:10], fmt).date()
        except ValueError:
            continue
    return None


def _parse_tipo(raw) -> Optional[models.MovementType]:
    if raw is None:
        return None
    key = _strip_accents(str(raw).strip().lower())
    if key in _ENTRADA:
        return models.MovementType.IN
    if key in _SAIDA:
        return models.MovementType.OUT
    return None


def _norm_tipo_banco(val) -> Optional[str]:
    """`movement_type` lido como texto: o banco convive com IN/OUT e Entrada/Saída."""
    if val is None:
        return None
    key = _strip_accents(str(val).strip().lower())
    if key in ("in", "entrada"):
        return "IN"
    if key in ("out", "saida"):
        return "OUT"
    return None


# ─────────────────────────────────────────────────────────
# VALIDAÇÃO (compartilhada pelo analyze e pelo lancar)
# ─────────────────────────────────────────────────────────

def _validate_rows(seller: models.Seller, rows: List[dict], db: Session) -> dict:
    """
    Valida as linhas de UM seller. `rows` vem com: line, movement_date,
    movement_type, sku, quantity, nf_number, observation.

    Devolve as linhas normalizadas (cada uma com `errors` e `warnings`), as
    listas gerais de erros/avisos e o resumo de saldo por SKU. Não grava nada.
    """
    hoje = today_brasilia()
    errors: List[str] = []
    warnings: List[str] = []
    out: List[dict] = []

    # ── 1. Campo a campo ─────────────────────────────────────────────────────
    for i, row in enumerate(rows):
        n = row.get("line") or (i + 2)
        r_err: List[str] = []

        data_raw = row.get("movement_date")
        data = _parse_data(data_raw)
        if data is None:
            if data_raw is None or str(data_raw).strip() == "":
                r_err.append("data não informada")
            else:
                r_err.append(f'data inválida ("{_cell_str(data_raw)}") — use dd/mm/aaaa')
        elif data > hoje:
            r_err.append(f"data {data.strftime('%d/%m/%Y')} está no futuro")

        tipo = _parse_tipo(row.get("movement_type"))
        if tipo is None:
            r_err.append(
                f'tipo inválido ("{_cell_str(row.get("movement_type"))}") — use Entrada, Saída, E ou S'
            )

        sku = _cell_str(row.get("sku"))
        if not sku:
            r_err.append("SKU não informado")

        qtd = _parse_qtd(row.get("quantity"))
        if qtd is None:
            r_err.append(
                f'quantidade inválida ("{_cell_str(row.get("quantity"))}") — '
                f"precisa ser um número inteiro maior que zero"
            )

        nf = _cell_str(row.get("nf_number"))
        obs = _cell_str(row.get("observation"))
        if not nf and not obs:
            r_err.append("preencha a NF ou a Observação (pelo menos uma)")
        if len(nf) > NF_MAX:
            r_err.append(f"NF com mais de {NF_MAX} caracteres")

        out.append({
            "line": n,
            "movement_date": data,
            "date_raw": _cell_str(data_raw) if data is None else None,
            "movement_type": tipo,
            "type_raw": _cell_str(row.get("movement_type")),
            "sku": sku,
            "sku_typed": sku,
            "product_id": None,
            "product_name": None,
            "quantity": qtd,
            "quantity_raw": _cell_str(row.get("quantity")),
            "nf_number": nf,
            "observation": obs,
            "errors": r_err,
            "warnings": [],
        })

    # ── 2. SKU contra o cadastro do seller (uma consulta, sem diferenciar caixa) ─
    # ⚠️ 28/08/2026: 'MOSQ2' e 'mosq2' são o mesmo produto — grava sempre com a
    # grafia do cadastro, senão parte o estoque em duas posições.
    lowers = {r["sku"].lower() for r in out if r["sku"]}
    produtos: dict = {}
    discontinuados: set = set()
    if lowers:
        for p in db.query(models.Product).filter(
            models.Product.seller_id == seller.id,
            models.Product.active == True,  # noqa: E712
            func.lower(models.Product.sku).in_(list(lowers)),
        ).order_by(models.Product.id).all():
            produtos.setdefault(p.sku.lower(), p)
        discontinuados = {
            d[0].lower() for d in db.query(models.DiscontinuedSku.sku).filter(
                models.DiscontinuedSku.seller_id == seller.id,
                func.lower(models.DiscontinuedSku.sku).in_(list(lowers)),
            ).all()
        }

    for r in out:
        if not r["sku"]:
            continue
        key = r["sku"].lower()
        prod = produtos.get(key)
        if key in discontinuados:
            r["errors"].append(f'SKU "{r["sku"]}" foi descontinuado e não aceita movimentação')
        elif not prod:
            r["errors"].append(f'SKU "{r["sku"]}" não tem produto ativo cadastrado neste seller')
        else:
            r["sku"] = prod.sku
            r["product_id"] = prod.id
            r["product_name"] = prod.name

    # ── 3. Avisos: linha idêntica no mesmo arquivo ──────────────────────────
    def _chave(r):
        return (
            r["movement_date"], r["movement_type"], r["sku"].lower(), r["quantity"],
            r["nf_number"].lower(), r["observation"].lower(),
        )

    completas = [r for r in out if not r["errors"]]
    vistos: dict = {}
    for r in completas:
        k = _chave(r)
        if k in vistos:
            r["warnings"].append(f"idêntica à linha {vistos[k]}")
        else:
            vistos[k] = r["line"]

    # ── 4. Avisos: linha que parece já lançada antes ─────────────────────────
    # Uma consulta só, pelo índice (seller_id, sku). `movement_type` vem como
    # texto (CAST) porque o banco convive com IN/OUT e Entrada/Saída — carregar
    # pelo ORM quebraria no valor legado do SQLite.
    if completas:
        skus = list({r["sku"] for r in completas})
        datas = list({r["movement_date"] for r in completas})
        sm = models.StockMovement
        existentes = set()
        for sku_db, dt_db, tipo_db, qtd_db, nf_db, obs_db in db.query(
            sm.sku, sm.movement_date, cast(sm.movement_type, String),
            sm.quantity, sm.nf_number, sm.observation,
        ).filter(
            sm.seller_id == seller.id,
            sm.sku.in_(skus),
            sm.movement_date.in_(datas),
        ).all():
            if isinstance(dt_db, str):  # SQLite devolve texto
                dt_db = _parse_data(dt_db)
            existentes.add((
                sku_db, dt_db, _norm_tipo_banco(tipo_db), qtd_db,
                (nf_db or "").strip().lower(), (obs_db or "").strip().lower(),
            ))
        for r in completas:
            k = (
                r["sku"], r["movement_date"], r["movement_type"].name, r["quantity"],
                r["nf_number"].lower(), r["observation"].lower(),
            )
            if k in existentes:
                r["warnings"].append("parece já lançada antes (já existe movimentação igual)")

    # ── 5. Saldo antes → depois por SKU ──────────────────────────────────────
    entradas: dict = defaultdict(int)
    saidas: dict = defaultdict(int)
    nomes: dict = {}
    for r in completas:
        if r["movement_type"] == models.MovementType.IN:
            entradas[r["sku"]] += r["quantity"]
        else:
            saidas[r["sku"]] += r["quantity"]
        nomes[r["sku"]] = r["product_name"]

    saldos: dict = {}
    if nomes:
        for sku_db, atual in db.query(
            models.StockPosition.sku, models.StockPosition.current_stock,
        ).filter(
            models.StockPosition.seller_id == seller.id,
            models.StockPosition.sku.in_(list(nomes)),
        ).all():
            saldos[sku_db] = atual or 0

    resumo = []
    for sku in sorted(nomes):
        antes = saldos.get(sku, 0)
        depois = antes + entradas[sku] - saidas[sku]
        resumo.append({
            "sku": sku,
            "product_name": nomes[sku],
            "stock_before": antes,
            "total_in": entradas[sku],
            "total_out": saidas[sku],
            "stock_after": depois,
            "goes_negative": depois < 0,
        })
        if depois < 0:
            warnings.append(
                f'SKU "{sku}" vai ficar com saldo negativo ({antes} → {depois})'
            )

    for r in out:
        for e in r["errors"]:
            errors.append(f"Linha {r['line']}: {e}")
        for w in r["warnings"]:
            warnings.append(f"Linha {r['line']}: {w}")

    return {"rows": out, "errors": errors, "warnings": warnings, "skus": resumo}


def _row_to_payload(r: dict) -> dict:
    """Formato devolvido à tela (e que a tela manda de volta no /lancar)."""
    return {
        "line": r["line"],
        "movement_date": r["movement_date"].isoformat() if r["movement_date"] else r["date_raw"],
        "movement_type": _TIPO_LABEL.get(r["movement_type"], r["type_raw"]),
        "sku": r["sku"],
        "product_name": r["product_name"],
        "quantity": r["quantity"] if r["quantity"] is not None else r["quantity_raw"],
        "nf_number": r["nf_number"],
        "observation": r["observation"],
        "errors": r["errors"],
        "warnings": r["warnings"],
    }


def _resultado(seller: models.Seller, v: dict) -> dict:
    rows = [_row_to_payload(r) for r in v["rows"]]
    return {
        "seller_id": seller.id,
        "seller_name": seller.trade_name or seller.name,
        "total": len(rows),
        "entradas": sum(1 for r in v["rows"] if r["movement_type"] == models.MovementType.IN),
        "saidas": sum(1 for r in v["rows"] if r["movement_type"] == models.MovementType.OUT),
        "rows": rows,
        "errors": v["errors"],
        "warnings": v["warnings"],
        "skus": v["skus"],
        "can_submit": len(v["errors"]) == 0,
    }


# ─────────────────────────────────────────────────────────
# ANÁLISE DA PLANILHA (não grava nada)
# ─────────────────────────────────────────────────────────

@router.post("/{seller_id}/analyze")
def analyze_file(
    seller_id: int,
    file: UploadFile = File(...),
    current_user: models.User = Depends(require_manager_or_above),
    user_seller_ids: Optional[List[int]] = Depends(get_user_seller_ids),
    db: Session = Depends(get_db),
):
    """
    Lê a planilha e devolve a conferência. NÃO grava nada. Síncrono de
    propósito (`def`): ler Excel é trabalho de CPU e travaria o event loop.
    """
    seller = _get_seller_in_scope(seller_id, user_seller_ids, db)

    nome = (file.filename or "").lower()
    if not nome.endswith(".xlsx"):
        raise HTTPException(status_code=422, detail="Envie um arquivo .xlsx (use o modelo)")

    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(file.file.read()), data_only=True, read_only=True)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Não consegui ler o arquivo: {exc}")

    ws = wb[ABA] if ABA in wb.sheetnames else wb.worksheets[0]

    # Linha 1 = cabeçalho. A numeração devolvida é a MESMA do Excel (a 1ª linha
    # de dados é a 2), pra pessoa achar a linha certa na planilha.
    brutas: List[dict] = []
    for excel_row, raw in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        celulas = list(raw) + [None] * 6
        if all(c is None or str(c).strip() == "" for c in celulas[:6]):
            continue
        brutas.append({
            "line": excel_row,
            "movement_date": celulas[0],
            "movement_type": celulas[1],
            "sku": celulas[2],
            "quantity": celulas[3],
            "nf_number": celulas[4],
            "observation": celulas[5],
        })
    wb.close()

    if not brutas:
        raise HTTPException(status_code=422, detail="Nenhuma linha preenchida na planilha")

    return _resultado(seller, _validate_rows(seller, brutas, db))


# ─────────────────────────────────────────────────────────
# LANÇAMENTO (linhas conferidas na tela)
# ─────────────────────────────────────────────────────────

@router.post("/{seller_id}/lancar", status_code=201)
def lancar(
    seller_id: int,
    body: dict,
    current_user: models.User = Depends(require_manager_or_above),
    user_seller_ids: Optional[List[int]] = Depends(get_user_seller_ids),
    db: Session = Depends(get_db),
):
    """
    Grava. TUDO-OU-NADA: revalida do zero (a tela pode ser burlada e um erro
    aqui vira estoque errado) e, havendo qualquer erro, devolve 422 sem gravar.
    Avisos não bloqueiam — a pessoa já os viu na conferência.
    """
    seller = _get_seller_in_scope(seller_id, user_seller_ids, db)

    rows = body.get("rows") or []
    if not rows:
        raise HTTPException(status_code=422, detail="Nenhuma linha enviada")

    v = _validate_rows(seller, rows, db)
    if v["errors"]:
        raise HTTPException(
            status_code=422,
            detail={
                "message": "Nada foi lançado — corrija os problemas abaixo.",
                "errors": v["errors"],
            },
        )

    agora = now_brasilia()
    entradas = saidas = 0
    delta: dict = defaultdict(lambda: {models.MovementType.IN: 0, models.MovementType.OUT: 0})
    nomes: dict = {}
    try:
        for r in v["rows"]:
            db.add(models.StockMovement(
                seller_id=seller.id,
                product_id=r["product_id"],
                sku=r["sku"],
                product_name=r["product_name"] or r["sku"],
                movement_date=r["movement_date"],
                movement_type=r["movement_type"],
                quantity=r["quantity"],
                adjusted_quantity=r["quantity"],
                nf_number=r["nf_number"] or None,
                nature=NATUREZA,
                # order_id fica VAZIO de propósito — ver o cabeçalho do módulo.
                observation=r["observation"] or None,
                operator_id=current_user.id,
                created_at=agora,
            ))
            delta[r["sku"]][r["movement_type"]] += r["quantity"]
            nomes[r["sku"]] = r["product_name"] or r["sku"]
            if r["movement_type"] == models.MovementType.IN:
                entradas += 1
            else:
                saidas += 1

        # Saldo: uma soma atômica por SKU e tipo, SEMPRE em ordem de SKU — dois
        # lotes simultâneos travam as posições na mesma ordem e não entram em
        # deadlock (mesmo cuidado do `apply_stock_for_orders`, 16/09/2026).
        for sku in sorted(delta):
            for mt, qtd in delta[sku].items():
                if qtd:
                    update_stock_position(
                        seller_id=seller.id,
                        sku=sku,
                        product_name=nomes[sku],
                        movement_type=mt,
                        quantity=qtd,
                        db=db,
                    )

        db.add(models.AuditLog(
            entity_type="StockMovement",
            entity_id=seller.id,
            action="BULK_UPLOAD",
            detail=(
                f"Lançamento de planilha | Seller: {seller.trade_name or seller.name} | "
                f"Linhas: {len(v['rows'])} (Entradas: {entradas} | Saídas: {saidas}) | "
                f"Avisos: {len(v['warnings'])}"
            ),
            user_id=current_user.id,
            timestamp=agora,
        ))
        db.commit()
    except Exception as exc:
        db.rollback()
        raise HTTPException(
            status_code=500,
            detail=f"Falha ao lançar — nada foi gravado: {exc}",
        )

    return {
        "total": len(v["rows"]),
        "entradas": entradas,
        "saidas": saidas,
        "warnings": len(v["warnings"]),
    }
