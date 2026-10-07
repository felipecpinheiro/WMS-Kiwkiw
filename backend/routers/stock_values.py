"""
WMS Kiwkiw - Valor unitário por SKU no Meu Estoque (Portal do Seller) — 07/10/2026

O PRÓPRIO seller (role `client`) cadastra o valor unitário (R$) de cada SKU. A
tela de Meu Estoque mostra Valor Un, Valor Total e o total geral. É só
cadastro e visualização: nada aqui toca em stock_movements, stock_positions,
pedidos ou faturamento.

  GET  /inventory/valores/modelo    -> Excel modelo em memória (1 linha por SKU)
  POST /inventory/valores/analyze   -> confere a planilha, NÃO grava nada
  POST /inventory/valores/lancar    -> revalida do zero e grava (tudo-ou-nada)
  PUT  /inventory/valores/sku       -> edição individual (valor ou apagar)

Onde mora o dado: `products.seller_unit_value` NUMERIC(12,2) NULL.
⚠️ NÃO é a `products.unit_value` legada (Float, default 0.0, editada nas telas
internas). NULL = "sem valor"; 0 é valor válido (brinde/amostra).

Permissões: só `client` grava, e só no próprio seller_id (sempre de
`current_user.seller_id`, nunca de parâmetro). `admin` só lê (modelo/analyze),
com `?seller_id=`. Manager e operator: 403.

Regras da planilha (tudo-ou-nada):
  * Bloqueiam: SKU sem produto ativo no seller (casa sem diferenciar caixa,
    grava a grafia do cadastro), SKU descontinuado, valor não numérico,
    negativo, mais de 2 casas, mesmo SKU repetido com valores diferentes.
  * Só avisam: célula de valor vazia (MANTÉM o atual, nunca apaga), SKU do
    seller ausente da planilha (não é alterado), SKU repetido com o mesmo valor.
  * Apagar valor só pelo site (PUT /sku com value=null).
"""

import io
import re
from decimal import Decimal, InvalidOperation
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException, File, Query, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from ..database import get_db
from ..auth import get_current_user
from ..timezone_utils import now_brasilia
from .. import models
from .returns import _cell_str

router = APIRouter(prefix="/inventory/valores", tags=["Estoque — Valor unitário do seller"])

ABA = "VALORES"
COLUNAS_MODELO = ["SKU", "Nome do produto", "Valor Un"]
_MAX = Decimal("9999999999.99")  # cabe em NUMERIC(12,2)


# ─────────────────────────────────────────────────────────
# ESCOPO
# ─────────────────────────────────────────────────────────

def _role(user: models.User) -> str:
    return user.role.value if hasattr(user.role, "value") else user.role


def _seller_for_write(user: models.User, db: Session) -> models.Seller:
    """Só `client`, só o próprio seller (ativo)."""
    if _role(user) != "client":
        raise HTTPException(status_code=403, detail="Acesso negado — só o seller cadastra valores")
    return _active_seller(user.seller_id, db)


def _seller_for_read(user: models.User, seller_id: Optional[int], db: Session) -> models.Seller:
    """client -> o próprio; admin -> exige ?seller_id= (leitura de suporte)."""
    role = _role(user)
    if role == "client":
        return _active_seller(user.seller_id, db)
    if role == "admin":
        if not seller_id:
            raise HTTPException(status_code=422, detail="Informe seller_id")
        return _active_seller(seller_id, db)
    raise HTTPException(status_code=403, detail="Acesso negado")


def _active_seller(seller_id: Optional[int], db: Session) -> models.Seller:
    if not seller_id:
        raise HTTPException(status_code=400, detail="Usuário sem seller associado")
    seller = db.query(models.Seller).filter(
        models.Seller.id == seller_id,
        models.Seller.active == True,  # noqa: E712
    ).first()
    if not seller:
        raise HTTPException(status_code=404, detail=f"Seller {seller_id} não encontrado ou inativo")
    return seller


# ─────────────────────────────────────────────────────────
# CONVERSÃO DO VALOR
# ─────────────────────────────────────────────────────────

def _parse_valor(raw):
    """
    Devolve (Decimal | None, erro | None).
    None sem erro = célula vazia. Aceita número do Excel, "12,50", "12.5",
    "R$ 12,50" e "1.234,56" (separador de milhar só quando os dois aparecem).
    """
    if raw is None:
        return None, None
    if isinstance(raw, bool):
        return None, "valor inválido"
    if isinstance(raw, (int, float, Decimal)):
        txt = str(raw)
    else:
        txt = str(raw).strip()
        if txt == "":
            return None, None
        txt = re.sub(r"^R\$\s*", "", txt, flags=re.IGNORECASE).replace(" ", "")
        if "," in txt and "." in txt:
            if txt.rfind(",") > txt.rfind("."):
                txt = txt.replace(".", "").replace(",", ".")
            else:
                txt = txt.replace(",", "")
        else:
            txt = txt.replace(",", ".")
    try:
        d = Decimal(txt)
    except InvalidOperation:
        return None, f'valor "{_cell_str(raw)}" não é um número'
    if not d.is_finite():
        return None, f'valor "{_cell_str(raw)}" não é um número'
    if d < 0:
        return None, "valor não pode ser negativo"
    d = d.normalize()
    if d.as_tuple().exponent < -2:
        return None, "valor com mais de 2 casas decimais"
    if d > _MAX:
        return None, "valor grande demais"
    return d.quantize(Decimal("0.01")), None


def _num(d: Optional[Decimal]) -> Optional[float]:
    return float(d) if d is not None else None


def _atual(p: models.Product) -> Optional[Decimal]:
    v = p.seller_unit_value
    return Decimal(str(v)).quantize(Decimal("0.01")) if v is not None else None


# ─────────────────────────────────────────────────────────
# CADASTRO DO SELLER
# ─────────────────────────────────────────────────────────

def _produtos_ativos(seller_id: int, db: Session) -> dict:
    """{sku em minúsculas: Product} dos produtos ativos NÃO descontinuados."""
    descont = {
        r[0].lower() for r in db.query(models.DiscontinuedSku.sku).filter(
            models.DiscontinuedSku.seller_id == seller_id,
        ).all()
    }
    prods: dict = {}
    for p in db.query(models.Product).filter(
        models.Product.seller_id == seller_id,
        models.Product.active == True,  # noqa: E712
    ).order_by(models.Product.id).all():
        if p.sku.lower() in descont:
            continue
        prods.setdefault(p.sku.lower(), p)
    return prods


def _descontinuados(seller_id: int, db: Session) -> set:
    return {
        r[0].lower() for r in db.query(models.DiscontinuedSku.sku).filter(
            models.DiscontinuedSku.seller_id == seller_id,
        ).all()
    }


# ─────────────────────────────────────────────────────────
# MODELO
# ─────────────────────────────────────────────────────────

@router.get("/modelo")
def download_modelo(
    seller_id: Optional[int] = Query(None),
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Excel modelo EM MEMÓRIA: 1 linha por SKU ativo (inclui saldo zero/negativo)."""
    seller = _seller_for_read(current_user, seller_id, db)

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = ABA
    ws.append(COLUNAS_MODELO)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7B63E8")
        cell.alignment = Alignment(horizontal="center")

    for key, p in sorted(_produtos_ativos(seller.id, db).items(), key=lambda kv: kv[1].sku):
        atual = _atual(p)
        ws.append([p.sku, p.name, float(atual) if atual is not None else None])
        ws.cell(row=ws.max_row, column=3).number_format = "0.00"
        ws.cell(row=ws.max_row, column=1).number_format = "@"

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 14
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("INSTRUCOES")
    for linha in [
        ["Como preencher"],
        [""],
        ["SKU", "Nao altere. Cada linha e um produto ativo seu."],
        ["Nome do produto", "Apenas para voce se localizar. Nao e lido."],
        ["Valor Un", "Valor unitario em reais, ate 2 casas (12,50 ou 12.5). Zero e permitido."],
        [""],
        ["Observacoes"],
        ["", "Valor Un em branco MANTEM o valor atual (nunca apaga). Para apagar, use o site."],
        ["", "Linhas que voce apagar da planilha nao sao alteradas."],
        ["", "Se qualquer linha tiver problema, NADA e gravado."],
    ]:
        ws2.append(linha)
    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 95

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="MODELO_VALORES_ESTOQUE.xlsx"'},
    )


# ─────────────────────────────────────────────────────────
# VALIDAÇÃO (compartilhada pelo analyze e pelo lancar)
# ─────────────────────────────────────────────────────────

def _validate_rows(seller: models.Seller, rows: List[dict], db: Session) -> dict:
    """
    `rows` vem com: line, sku, value (cru). Não grava nada. Devolve as linhas
    normalizadas (com errors/warnings), erros/avisos gerais e o antes→depois.
    """
    errors: List[str] = []
    warnings: List[str] = []
    out: List[dict] = []

    ativos = _produtos_ativos(seller.id, db)
    descont = _descontinuados(seller.id, db)

    for i, row in enumerate(rows):
        n = row.get("line") or (i + 2)
        r_err: List[str] = []
        r_warn: List[str] = []

        sku_typed = _cell_str(row.get("sku"))
        prod = None
        if not sku_typed:
            r_err.append("SKU não informado")
        else:
            key = sku_typed.lower()
            if key in descont:
                r_err.append(f'SKU "{sku_typed}" foi descontinuado')
            elif key not in ativos:
                r_err.append(f'SKU "{sku_typed}" não tem produto ativo cadastrado neste seller')
            else:
                prod = ativos[key]

        novo, erro = _parse_valor(row.get("value"))
        if erro:
            r_err.append(erro)
        elif novo is None and prod is not None:
            r_warn.append("valor em branco — mantém o valor atual")

        out.append({
            "line": n,
            "sku": prod.sku if prod else sku_typed,
            "product_id": prod.id if prod else None,
            "product_name": prod.name if prod else None,
            "current_value": _atual(prod) if prod else None,
            "new_value": novo,
            "errors": r_err,
            "warnings": r_warn,
        })

    # Mesmo SKU repetido (sem diferenciar caixa): valores diferentes bloqueiam.
    por_sku: dict = {}
    for r in out:
        if r["product_id"] is not None and not r["errors"]:
            por_sku.setdefault(r["sku"].lower(), []).append(r)
    for lista in por_sku.values():
        if len(lista) < 2:
            continue
        com_valor = [r for r in lista if r["new_value"] is not None]
        distintos = {r["new_value"] for r in com_valor}
        primeira = lista[0]["line"]
        if len(distintos) > 1:
            for r in lista[1:]:
                r["errors"].append(
                    f"SKU repetido com valores diferentes (já aparece na linha {primeira})")
        else:
            for r in lista[1:]:
                r["warnings"].append(f"SKU repetido (já aparece na linha {primeira})")

    # Antes → depois (só linhas válidas com valor novo diferente do atual).
    vistos: set = set()
    mudancas = []
    for r in out:
        if r["errors"] or r["product_id"] is None or r["new_value"] is None:
            continue
        k = r["sku"].lower()
        if k in vistos:
            continue
        vistos.add(k)
        if r["current_value"] != r["new_value"]:
            mudancas.append({
                "sku": r["sku"],
                "product_name": r["product_name"],
                "before": _num(r["current_value"]),
                "after": _num(r["new_value"]),
            })

    # SKU do seller ausente da planilha: 1 aviso agregado.
    presentes = {r["sku"].lower() for r in out if r["product_id"] is not None}
    ausentes = len([k for k in ativos if k not in presentes])
    if ausentes:
        warnings.append(f"{ausentes} SKU(s) do seu cadastro não estão na planilha e não serão alterados")

    for r in out:
        for e in r["errors"]:
            errors.append(f"Linha {r['line']}: {e}")
        for w in r["warnings"]:
            warnings.append(f"Linha {r['line']}: {w}")

    return {"rows": out, "errors": errors, "warnings": warnings, "changes": mudancas}


def _row_payload(r: dict) -> dict:
    """Formato devolvido à tela (e que ela manda de volta no /lancar)."""
    return {
        "line": r["line"],
        "sku": r["sku"],
        "product_name": r["product_name"],
        "current_value": _num(r["current_value"]),
        "value": _num(r["new_value"]),
        "errors": r["errors"],
        "warnings": r["warnings"],
    }


# ─────────────────────────────────────────────────────────
# ANÁLISE (não grava nada)
# ─────────────────────────────────────────────────────────

@router.post("/analyze")
def analyze_file(
    seller_id: Optional[int] = Query(None),
    file: UploadFile = File(...),
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Lê a planilha e devolve a conferência. NÃO grava. Síncrono (`def`) de propósito."""
    seller = _seller_for_read(current_user, seller_id, db)

    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(status_code=422, detail="Envie um arquivo .xlsx (use o modelo)")

    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(file.file.read()), data_only=True, read_only=True)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Não consegui ler o arquivo: {exc}")

    ws = wb[ABA] if ABA in wb.sheetnames else wb.worksheets[0]

    # 1ª linha de dados = linha 2 do Excel (mesma numeração mostrada na tela).
    brutas: List[dict] = []
    for excel_row, raw in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        c = list(raw) + [None] * 3
        if all(x is None or str(x).strip() == "" for x in c[:3]):
            continue
        brutas.append({"line": excel_row, "sku": c[0], "value": c[2]})
    wb.close()

    if not brutas:
        raise HTTPException(status_code=422, detail="Nenhuma linha preenchida na planilha")

    v = _validate_rows(seller, brutas, db)
    return {
        "seller_id": seller.id,
        "total": len(v["rows"]),
        "rows": [_row_payload(r) for r in v["rows"]],
        "errors": v["errors"],
        "warnings": v["warnings"],
        "changes": v["changes"],
        "can_submit": len(v["errors"]) == 0,
    }


# ─────────────────────────────────────────────────────────
# LANÇAMENTO
# ─────────────────────────────────────────────────────────

def _fmt(v: Optional[Decimal]) -> str:
    return "sem valor" if v is None else f"{v:.2f}"


@router.post("/lancar", status_code=201)
def lancar(
    body: dict,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """
    Grava. TUDO-OU-NADA: revalida do zero com a mesma função do analyze (a tela
    pode ser burlada). Só grava — e só loga — o que realmente mudou.
    """
    seller = _seller_for_write(current_user, db)

    rows = body.get("rows") or []
    if not rows:
        raise HTTPException(status_code=422, detail="Nenhuma linha enviada")

    v = _validate_rows(seller, rows, db)
    if v["errors"]:
        raise HTTPException(
            status_code=422,
            detail={"message": "Nada foi gravado — corrija os problemas abaixo.", "errors": v["errors"]},
        )

    agora = now_brasilia()
    try:
        ids = {r["product_id"] for r in v["rows"] if r["product_id"] is not None}
        produtos = {
            p.id: p for p in db.query(models.Product).filter(models.Product.id.in_(ids)).all()
        } if ids else {}

        feitos: set = set()
        alterados = []
        for r in v["rows"]:
            if r["new_value"] is None or r["product_id"] in feitos:
                continue
            feitos.add(r["product_id"])
            p = produtos[r["product_id"]]
            antes = _atual(p)
            if antes == r["new_value"]:
                continue
            p.seller_unit_value = r["new_value"]
            alterados.append(f"{p.sku}: {_fmt(antes)} -> {_fmt(r['new_value'])}")

        if alterados:
            db.add(models.AuditLog(
                entity_type="Product",
                entity_id=seller.id,
                action="BULK_UPLOAD",
                detail=(
                    f"Valor unitário por planilha | Seller: {seller.trade_name or seller.name} | "
                    f"Alterados: {len(alterados)} | " + "; ".join(alterados)
                ),
                user_id=current_user.id,
                timestamp=agora,
            ))
        db.commit()
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Falha ao gravar — nada foi gravado: {exc}")

    return {"total": len(v["rows"]), "changed": len(alterados), "warnings": len(v["warnings"])}


# ─────────────────────────────────────────────────────────
# EDIÇÃO INDIVIDUAL (inline na tabela)
# ─────────────────────────────────────────────────────────

@router.put("/sku")
def set_sku_value(
    body: dict,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """
    body = {"sku": "...", "value": 12.5 | "12,50" | null}. `null`/vazio APAGA o
    valor (volta a "sem valor") — só vale pelo site. SKU no corpo (e não na
    URL) porque SKU pode ter caracteres que quebram o caminho.
    """
    seller = _seller_for_write(current_user, db)

    sku_typed = _cell_str(body.get("sku"))
    if not sku_typed:
        raise HTTPException(status_code=422, detail="SKU não informado")
    key = sku_typed.lower()
    if key in _descontinuados(seller.id, db):
        raise HTTPException(status_code=422, detail=f'SKU "{sku_typed}" foi descontinuado')
    prod = _produtos_ativos(seller.id, db).get(key)
    if not prod:
        raise HTTPException(
            status_code=422, detail=f'SKU "{sku_typed}" não tem produto ativo cadastrado neste seller')

    novo, erro = _parse_valor(body.get("value"))
    if erro:
        raise HTTPException(status_code=422, detail=erro)

    antes = _atual(prod)
    if antes != novo:
        prod.seller_unit_value = novo
        db.add(models.AuditLog(
            entity_type="Product",
            entity_id=prod.id,
            action="UPDATE",
            detail=f"Valor unitário | Seller: {seller.trade_name or seller.name} | "
                   f"{prod.sku}: {_fmt(antes)} -> {_fmt(novo)}",
            user_id=current_user.id,
            timestamp=now_brasilia(),
        ))
        db.commit()

    return {"sku": prod.sku, "seller_unit_value": _num(novo)}
